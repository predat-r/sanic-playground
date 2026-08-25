import asyncio
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.slide import Slide
from pptx.util import Inches, Pt

from app.presentations.schemas import PresentationContent, SlideContent


_PPTX_EXECUTOR = ThreadPoolExecutor(max_workers=2, thread_name_prefix="pptx-render")


class PptxRenderError(RuntimeError):
    """Raised when presentation content cannot be rendered as a PPTX file."""


class PptxRenderer:
    """Render validated presentation content with python-pptx."""

    _NAVY = RGBColor(18, 32, 56)
    _BLUE = RGBColor(35, 100, 190)
    _LIGHT_BLUE = RGBColor(232, 240, 250)
    _LIGHT_GREY = RGBColor(244, 246, 248)
    _MID_GREY = RGBColor(96, 106, 120)
    _WHITE = RGBColor(255, 255, 255)

    async def render(self, content: PresentationContent) -> bytes:
        """Build the PPTX off the event loop because python-pptx is synchronous."""
        loop = asyncio.get_running_loop()
        return await loop.run_in_executor(
            _PPTX_EXECUTOR,
            self._render_blocking,
            content,
        )

    def _render_blocking(self, content: PresentationContent) -> bytes:
        try:
            presentation = Presentation()
            presentation.slide_width = Inches(13.333)
            presentation.slide_height = Inches(7.5)

            self._add_title_slide(presentation, content)
            self._add_summary_slide(presentation, content)
            for number, slide_content in enumerate(content.slides, start=3):
                self._add_content_slide(
                    presentation,
                    slide_content,
                    number=number,
                    disclaimer=content.disclaimer,
                )

            output = BytesIO()
            presentation.save(output)
            return output.getvalue()
        except PptxRenderError:
            raise
        except Exception as exc:
            raise PptxRenderError("Could not render the presentation") from exc

    def _add_title_slide(
        self,
        presentation: Presentation,
        content: PresentationContent,
    ) -> None:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        self._set_background(slide, self._NAVY)
        self._add_text(
            slide,
            content.deck_title,
            left=0.9,
            top=1.65,
            width=11.5,
            height=1.6,
            font_size=50,
            color=self._WHITE,
            bold=True,
        )
        self._add_text(
            slide,
            content.deck_subtitle,
            left=0.95,
            top=3.45,
            width=10.7,
            height=0.9,
            font_size=24,
            color=RGBColor(205, 219, 239),
        )
        self._add_accent_rule(slide, top=5.85, color=self._BLUE)
        self._add_text(
            slide,
            "Financial presentation",
            left=0.95,
            top=6.08,
            width=4.0,
            height=0.35,
            font_size=16,
            color=RGBColor(172, 188, 211),
        )

    def _add_summary_slide(
        self,
        presentation: Presentation,
        content: PresentationContent,
    ) -> None:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        self._set_background(slide, self._WHITE)
        self._add_slide_title(slide, "Executive summary")

        panel = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.85),
            Inches(1.65),
            Inches(11.65),
            Inches(4.65),
        )
        panel.fill.solid()
        panel.fill.fore_color.rgb = self._LIGHT_GREY
        panel.line.fill.background()
        self._add_text(
            slide,
            self._truncate(content.executive_summary, 1_400),
            left=1.25,
            top=2.0,
            width=10.85,
            height=3.9,
            font_size=22,
            color=self._NAVY,
        )
        self._add_footer(slide, 2, content.disclaimer)

    def _add_content_slide(
        self,
        presentation: Presentation,
        content: SlideContent,
        *,
        number: int,
        disclaimer: str,
    ) -> None:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        self._set_background(slide, self._WHITE)
        self._add_slide_title(slide, self._truncate(content.title, 105))
        self._add_text(
            slide,
            self._truncate(content.key_message, 280),
            left=0.85,
            top=1.35,
            width=7.35,
            height=0.85,
            font_size=24,
            color=self._BLUE,
            bold=True,
        )

        bullet_text = "\n".join(
            f"• {self._truncate(bullet, 220)}" for bullet in content.bullets
        )
        self._add_text(
            slide,
            bullet_text,
            left=0.85,
            top=2.35,
            width=7.25,
            height=3.9,
            font_size=18,
            color=self._NAVY,
            paragraph_spacing=9,
        )

        panel = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(8.55),
            Inches(1.45),
            Inches(3.9),
            Inches(4.85),
        )
        panel.fill.solid()
        panel.fill.fore_color.rgb = self._LIGHT_BLUE
        panel.line.fill.background()
        self._add_text(
            slide,
            content.visual.type.upper(),
            left=8.95,
            top=1.85,
            width=3.1,
            height=0.35,
            font_size=16,
            color=self._BLUE,
            bold=True,
        )
        self._add_text(
            slide,
            self._truncate(content.visual.title, 115),
            left=8.95,
            top=2.35,
            width=3.05,
            height=1.05,
            font_size=24,
            color=self._NAVY,
            bold=True,
        )
        self._add_text(
            slide,
            self._truncate(content.visual.description, 460),
            left=8.95,
            top=3.55,
            width=3.05,
            height=2.2,
            font_size=16,
            color=self._MID_GREY,
        )
        self._add_notes(slide, content)
        self._add_footer(slide, number, disclaimer)

    def _add_slide_title(self, slide: Slide, title: str) -> None:
        self._add_text(
            slide,
            title,
            left=0.85,
            top=0.48,
            width=11.65,
            height=0.62,
            font_size=35,
            color=self._NAVY,
            bold=True,
        )
        self._add_accent_rule(slide, top=1.16, color=self._BLUE)

    def _add_accent_rule(self, slide: Slide, *, top: float, color: RGBColor) -> None:
        rule = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.85),
            Inches(top),
            Inches(1.15),
            Inches(0.06),
        )
        rule.fill.solid()
        rule.fill.fore_color.rgb = color
        rule.line.fill.background()

    def _add_footer(self, slide: Slide, number: int, disclaimer: str) -> None:
        self._add_text(
            slide,
            self._truncate(disclaimer, 180),
            left=0.85,
            top=6.82,
            width=10.5,
            height=0.28,
            font_size=10,
            color=self._MID_GREY,
        )
        self._add_text(
            slide,
            str(number),
            left=11.8,
            top=6.78,
            width=0.6,
            height=0.3,
            font_size=11,
            color=self._MID_GREY,
            alignment=PP_ALIGN.RIGHT,
        )

    def _add_notes(self, slide: Slide, content: SlideContent) -> None:
        notes = content.speaker_notes.strip()
        if content.source_references:
            sources = "\n".join(f"- {source}" for source in content.source_references)
            notes = f"{notes}\n\n[Sources]\n{sources}" if notes else f"[Sources]\n{sources}"

        notes_frame = slide.notes_slide.notes_text_frame
        if notes_frame is not None:
            notes_frame.text = notes

    def _add_text(
        self,
        slide: Slide,
        text: str,
        *,
        left: float,
        top: float,
        width: float,
        height: float,
        font_size: int,
        color: RGBColor,
        bold: bool = False,
        alignment: PP_ALIGN = PP_ALIGN.LEFT,
        paragraph_spacing: int = 0,
    ) -> None:
        shape = slide.shapes.add_textbox(
            Inches(left),
            Inches(top),
            Inches(width),
            Inches(height),
        )
        frame = shape.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.margin_left = 0
        frame.margin_right = 0
        frame.margin_top = 0
        frame.margin_bottom = 0
        frame.vertical_anchor = MSO_ANCHOR.TOP

        for index, line in enumerate(text.splitlines() or [""]):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.alignment = alignment
            paragraph.space_after = Pt(paragraph_spacing)
            run = paragraph.add_run()
            run.text = line
            run.font.name = "Aptos"
            run.font.size = Pt(font_size)
            run.font.bold = bold
            run.font.color.rgb = color

    @staticmethod
    def _set_background(slide: Slide, color: RGBColor) -> None:
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = color

    @staticmethod
    def _truncate(text: str, limit: int) -> str:
        normalized = " ".join(text.split())
        if len(normalized) <= limit:
            return normalized
        return f"{normalized[: limit - 1].rstrip()}…"
